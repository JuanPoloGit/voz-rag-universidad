"""Extrae los tres .docx a JSON estructurado. Nada se escribe a mano."""
import json, re, sys
from pathlib import Path
import docx

ORIGEN = Path(sys.argv[1]); DESTINO = Path(sys.argv[2])

# --- Informe completo: secciones por Heading ------------------------------
d = docx.Document(str(ORIGEN / "162ac7b9-Informe_Completo_-_AudacIA_y_Universidad_Sim_n_Bol_var.docx"))
secciones, h1, h2 = {}, None, None
for p in d.paragraphs:
    txt = p.text.strip()
    if not txt: continue
    estilo = p.style.name
    if estilo == "Heading 1":
        h1, h2 = re.sub(r"^[^\w]*\d+\.\s*", "", txt).strip(), None
        secciones[h1] = {"_": []}
    elif estilo == "Heading 2":
        h2 = txt
        secciones.setdefault(h1, {"_": []})[h2] = []
    elif h1:
        secciones[h1][h2 or "_"].append(txt)

# --- Links: prosa por slug -------------------------------------------------
d2 = docx.Document(str(ORIGEN / "3e096684-Links_de_todos_los_proyectos.docx"))
prosa, slug = {}, None
for p in d2.paragraphs:
    t = p.text.strip()
    if not t: continue
    if t.startswith("http"):
        slug = t.rstrip("/").rsplit("/", 1)[-1]; prosa[slug] = []
    elif slug:
        prosa[slug].append(t)

# --- Universidad -----------------------------------------------------------
d3 = docx.Document(str(ORIGEN / "53a122f2-INFORME_UNIVERSIDAD_SIM_N_BOL_VAR.docx"))
uni = [p.text.strip() for p in d3.paragraphs if p.text.strip()]
tabla = [[c.text.strip() for c in fila.cells] for t in d3.tables for fila in t.rows]

DESTINO.write_text(json.dumps(
    {"secciones": secciones, "prosa": prosa, "universidad": uni, "programas": tabla},
    ensure_ascii=False, indent=1), encoding="utf-8")
print("secciones:", list(secciones))
print("proyectos con prosa:", len(prosa))
print("filas de programas:", len(tabla))
