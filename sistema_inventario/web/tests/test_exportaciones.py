"""
Pruebas de los servicios de exportación a Excel y CSV (web/servicios/excel_export.py y csv_export.py).
"""

from io import BytesIO
import openpyxl
from django.test import TestCase

from cuentas.models import UsuarioAudacIA
from inventario.models import Almacen, Categoria, Existencia, Movimiento, Prestamo, Producto
from web.servicios.csv_export import (
    exportar_existencias_csv,
    exportar_movimientos_csv,
    exportar_prestamos_csv,
    exportar_resumen_dashboard_csv,
)
from web.servicios.excel_export import (
    exportar_dashboard_excel,
    exportar_existencias_excel,
    exportar_movimientos_excel,
    exportar_prestamos_excel,
)
from web.servicios.reportes import (
    obtener_metricas_dashboard,
    obtener_reporte_existencias,
    obtener_reporte_movimientos_periodo,
    obtener_reporte_prestamos_activos,
)


class ServiciosExportacionTest(TestCase):
    def setUp(self):
        self.almacen = Almacen.objects.create(nombre="Almacén Pruebas")
        self.categoria = Categoria.objects.create(nombre="Electrónica", almacen_principal=self.almacen)

        self.admin = UsuarioAudacIA.objects.create_user(
            username="admin_exp_test",
            nombre_completo="Admin Export Test",
            email="admin_exp@audacia.edu.co",
            password="Password123!",
            rol=UsuarioAudacIA.Rol.ADMINISTRADOR,
        )

        self.producto = Producto.objects.create(
            codigo_interno="PROD-EXP-001",
            nombre="Osciloscopio Digital",
            categoria=self.categoria,
            stock_minimo=5,
        )

        self.existencia = Existencia.objects.create(
            producto=self.producto,
            almacen=self.almacen,
            stock_disponible=10,
            stock_en_prestamo=2,
        )

        self.movimiento = Movimiento.objects.create(
            clave_idempotencia="IDEM-TEST-EXP-001",
            tipo=Movimiento.Tipo.ENTRADA,
            producto=self.producto,
            almacen_destino=self.almacen,
            cantidad=10,
            stock_disponible_anterior=0,
            stock_disponible_posterior=10,
            usuario=self.admin,
            motivo="Dotación inicial",
        )

        self.prestamo = Prestamo.objects.create(
            producto=self.producto,
            almacen=self.almacen,
            movimiento=self.movimiento,
            usuario=self.admin,
            cantidad_original=2,
            cantidad_pendiente_devolucion=2,
            destinatario="Profesor Investigador",
            estado=Prestamo.Estado.ABIERTO,
        )

    def test_exportar_dashboard_excel_profesional(self):
        buffer = exportar_dashboard_excel(self.admin, almacen=self.almacen)
        self.assertIsInstance(buffer, BytesIO)
        self.assertGreater(len(buffer.getvalue()), 0)

        wb = openpyxl.load_workbook(buffer)
        nombres_hojas = wb.sheetnames
        self.assertEqual(len(nombres_hojas), 4)
        self.assertIn("Resumen Ejecutivo", nombres_hojas)
        self.assertIn("Existencias y Stock", nombres_hojas)
        self.assertIn("Últimos Movimientos", nombres_hojas)
        self.assertIn("Préstamos Pendientes", nombres_hojas)

        # Verificar hoja de existencias
        ws_ex = wb["Existencias y Stock"]
        self.assertEqual(ws_ex.cell(row=4, column=1).value, "Código Interno")
        self.assertEqual(ws_ex.cell(row=5, column=1).value, "PROD-EXP-001")
        self.assertEqual(ws_ex.cell(row=5, column=5).value, 10)

    def test_exportar_hojas_individuales_excel(self):
        # Existencias
        qs_ex = obtener_reporte_existencias(self.admin)
        buf_ex = exportar_existencias_excel(qs_ex)
        wb_ex = openpyxl.load_workbook(buf_ex)
        self.assertEqual(wb_ex.active.title, "Existencias")

        # Movimientos
        qs_mov = obtener_reporte_movimientos_periodo(self.admin)
        buf_mov = exportar_movimientos_excel(qs_mov)
        wb_mov = openpyxl.load_workbook(buf_mov)
        self.assertEqual(wb_mov.active.title, "Kardex de Movimientos")

        # Préstamos
        qs_pres = obtener_reporte_prestamos_activos(self.admin)
        buf_pres = exportar_prestamos_excel(qs_pres)
        wb_pres = openpyxl.load_workbook(buf_pres)
        self.assertEqual(wb_pres.active.title, "Préstamos")

    def test_exportar_csv_con_bom(self):
        # 1. Existencias CSV
        qs_ex = obtener_reporte_existencias(self.admin)
        resp_ex = exportar_existencias_csv(qs_ex)
        self.assertEqual(resp_ex.status_code, 200)
        self.assertIn("text/csv", resp_ex["Content-Type"])
        self.assertTrue(resp_ex.content.startswith(b"\xef\xbb\xbf"))  # UTF-8 BOM
        self.assertIn(b"PROD-EXP-001", resp_ex.content)

        # 2. Movimientos CSV
        qs_mov = obtener_reporte_movimientos_periodo(self.admin)
        resp_mov = exportar_movimientos_csv(qs_mov)
        self.assertTrue(resp_mov.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"Dotaci\xc3\xb3n inicial", resp_mov.content)

        # 3. Préstamos CSV
        qs_pres = obtener_reporte_prestamos_activos(self.admin)
        resp_pres = exportar_prestamos_csv(qs_pres)
        self.assertTrue(resp_pres.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"Profesor Investigador", resp_pres.content)

        # 4. Resumen Dashboard CSV
        metricas = obtener_metricas_dashboard(self.admin)
        resp_dash = exportar_resumen_dashboard_csv(metricas, usuario=self.admin)
        self.assertTrue(resp_dash.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"Total Productos en Cat\xc3\xa1logo", resp_dash.content)
