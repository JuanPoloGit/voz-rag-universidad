"""
Servicio de exportación a Excel profesional (.xlsx) con openpyxl.
Paleta corporativa AudacIA, diseño multi-hoja dinámico, fórmulas nativas y auto-ajuste de columnas.
"""

from datetime import datetime
from io import BytesIO
from typing import Any, List, Optional

from django.db.models import F
from django.utils import timezone
import openpyxl
from openpyxl.chart import BarChart, DoughnutChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from inventario.models import Almacen, Categoria, Existencia, Movimiento, Prestamo, Producto
from web.servicios.reportes import (
    obtener_datos_grafico_categorias,
    obtener_metricas_dashboard,
    obtener_reporte_existencias,
    obtener_reporte_movimientos_periodo,
    obtener_reporte_prestamos_activos,
)

# ═════════════════════════════════════════════════════════════════════════════
# ESTILOS CORPORATIVOS AUDACIA
# ═════════════════════════════════════════════════════════════════════════════

FONT_TITULO = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
FONT_SUBTITULO = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
FONT_SECCION = Font(name="Calibri", size=11, bold=True, color="1E293B")
FONT_CABECERA = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_DATOS = Font(name="Calibri", size=10, color="0F172A")
FONT_DATOS_BOLD = Font(name="Calibri", size=10, bold=True, color="0F172A")
FONT_DATOS_MUTED = Font(name="Calibri", size=9, color="64748B")
FONT_TOTAL = Font(name="Calibri", size=10, bold=True, color="0F172A")

FONT_KPI_LABEL = Font(name="Calibri", size=8, bold=True, color="64748B")
FONT_KPI_VALOR = Font(name="Calibri", size=16, bold=True, color="1E293B")
FONT_KPI_VALOR_WARN = Font(name="Calibri", size=16, bold=True, color="B45309")
FONT_KPI_VALOR_DANGER = Font(name="Calibri", size=16, bold=True, color="B91C1C")

FILL_BANNER = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
FILL_CABECERA = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
FILL_FILA_PAR = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_FILA_IMPAR = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_KPI_HEADER = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
FILL_KPI_BODY = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_KPI = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

BORDER_TOTAL = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="double", color="0F172A"),
)

BORDER_KPI = Border(
    left=Side(style="medium", color="CBD5E1"),
    right=Side(style="medium", color="CBD5E1"),
    top=Side(style="medium", color="CBD5E1"),
    bottom=Side(style="medium", color="CBD5E1"),
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def estilar_rango(ws, start_row: int, start_col: int, end_row: int, end_col: int, border=None, fill=None):
    """Garantiza que todas las celdas en un rango (incluso combinado) tengan borde y relleno."""
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            if border:
                cell.border = border
            if fill:
                cell.fill = fill


def autoajustar_columnas(ws, min_width: int = 12, max_width: int = 50):
    """Ajusta automáticamente el ancho de las columnas según su contenido."""
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if val.startswith("="):
                max_len = max(max_len, 12)
            else:
                max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


def aplicar_estilo_fila(ws, row_idx: int, valores: List[Any], es_cabecera: bool = False, es_par: bool = False, alineaciones: Optional[List[str]] = None, formatos: Optional[List[str]] = None, es_total: bool = False):
    """Aplica formato corporativo uniforme a una fila de datos."""
    for col_idx, val in enumerate(valores, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)

        if es_cabecera:
            cell.font = FONT_CABECERA
            cell.fill = FILL_CABECERA
            cell.border = BORDER_THIN
            cell.alignment = ALIGN_CENTER
        elif es_total:
            cell.font = FONT_TOTAL
            cell.fill = FILL_TOTAL
            cell.border = BORDER_TOTAL
            cell.alignment = ALIGN_RIGHT if (alineaciones and col_idx <= len(alineaciones) and alineaciones[col_idx - 1] == "right") else ALIGN_LEFT
        else:
            cell.font = FONT_DATOS
            cell.fill = FILL_FILA_PAR if es_par else FILL_FILA_IMPAR
            cell.border = BORDER_THIN

            align_type = alineaciones[col_idx - 1] if alineaciones and col_idx <= len(alineaciones) else "left"
            if align_type == "center":
                cell.alignment = ALIGN_CENTER
            elif align_type == "right":
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT

        if formatos and col_idx <= len(formatos) and formatos[col_idx - 1]:
            cell.number_format = formatos[col_idx - 1]


def crear_banner_hoja(ws, titulo: str, subtitulo: str, total_columnas: int):
    """Crea el bloque superior corporativo con título y metadatos."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columnas)
    cell_t = ws.cell(row=1, column=1, value=f"  {titulo.upper()}")
    cell_t.font = FONT_TITULO
    cell_t.fill = FILL_BANNER
    cell_t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columnas)
    cell_s = ws.cell(row=2, column=1, value=f"  {subtitulo}")
    cell_s.font = FONT_SUBTITULO
    cell_s.fill = FILL_BANNER
    cell_s.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # Línea divisoria en blanco
    ws.row_dimensions[3].height = 8


# ═════════════════════════════════════════════════════════════════════════════
# EXPORTACIÓN DEL DASHBOARD COMPLETO (LIBRO MULTIDISCO PROFESIONAL)
# ═════════════════════════════════════════════════════════════════════════════

def exportar_dashboard_excel(usuario: Any, almacen: Optional[Almacen] = None) -> BytesIO:
    """
    Genera el libro Excel dinámico y profesional del Dashboard de AudacIA.
    Contiene 4 pestañas integrales con fórmulas y estilos corporativos.
    """
    wb = openpyxl.Workbook()
    ahora_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    nombre_almacen = almacen.nombre if almacen else "Todos los Almacenes (Global)"
    metricas = obtener_metricas_dashboard(usuario, almacen=almacen)

    # ─────────────────────────────────────────────────────────────────────────
    # PESTAÑA 1: RESUMEN EJECUTIVO
    # ─────────────────────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Resumen Ejecutivo"
    crear_banner_hoja(ws1, "AUDACIA — SISTEMA DE GESTIÓN DE INVENTARIO", "INFORME EJECUTIVO Y TABLERO DE CONTROL", 6)

    # Metadatos del informe
    meta_filas = [
        ("Generado por:", getattr(usuario, "nombre_completo", "") or usuario.username),
        ("Fecha y corte:", ahora_str),
        ("Ámbito de almacén:", nombre_almacen),
    ]
    r_meta = 4
    for etiqueta, valor in meta_filas:
        c1 = ws1.cell(row=r_meta, column=1, value=etiqueta)
        c1.font = FONT_DATOS_BOLD
        c2 = ws1.cell(row=r_meta, column=2, value=valor)
        c2.font = FONT_DATOS
        r_meta += 1

    ws1.row_dimensions[r_meta].height = 10
    r_kpi = r_meta + 1

    # Título KPIs
    ws1.merge_cells(start_row=r_kpi, start_column=1, end_row=r_kpi, end_column=6)
    c_kpi_t = ws1.cell(row=r_kpi, column=1, value="INDICADORES CLAVE DE RENDIMIENTO (KPIs)")
    c_kpi_t.font = FONT_SECCION
    r_kpi += 1

    # Bloques de KPIs (2 filas de 3 tarjetas)
    kpis = [
        ("TOTAL PRODUCTOS", metricas["total_productos"], FONT_KPI_VALOR),
        ("STOCK DISPONIBLE", metricas["total_stock_disponible"], FONT_KPI_VALOR),
        ("EN PRÉSTAMO", metricas["total_stock_prestamo"], FONT_KPI_VALOR),
        ("STOCK CRÍTICO (0)", metricas["total_criticos"], FONT_KPI_VALOR_DANGER),
        ("BAJO STOCK MÍNIMO", metricas["total_bajo_stock"], FONT_KPI_VALOR_WARN),
        ("PRÉSTAMOS ACTIVOS", metricas["total_prestamos_activos"], FONT_KPI_VALOR),
    ]

    col_pos = [1, 3, 5]
    for idx, (label, val, font_v) in enumerate(kpis):
        fila_base = r_kpi if idx < 3 else r_kpi + 3
        c_idx = col_pos[idx % 3]

        ws1.row_dimensions[fila_base].height = 18
        ws1.row_dimensions[fila_base + 1].height = 26

        # Fila label
        ws1.merge_cells(start_row=fila_base, start_column=c_idx, end_row=fila_base, end_column=c_idx + 1)
        estilar_rango(ws1, fila_base, c_idx, fila_base, c_idx + 1, border=BORDER_THIN, fill=FILL_KPI_HEADER)
        cl = ws1.cell(row=fila_base, column=c_idx, value=label)
        cl.font = FONT_KPI_LABEL
        cl.alignment = ALIGN_CENTER

        # Fila valor
        ws1.merge_cells(start_row=fila_base + 1, start_column=c_idx, end_row=fila_base + 1, end_column=c_idx + 1)
        estilar_rango(ws1, fila_base + 1, c_idx, fila_base + 1, c_idx + 1, border=BORDER_THIN, fill=FILL_KPI_BODY)
        cv = ws1.cell(row=fila_base + 1, column=c_idx, value=val)
        cv.font = font_v
        cv.alignment = ALIGN_CENTER
        cv.number_format = "#,##0"

    r_cat = r_kpi + 7
    datos_cat = obtener_datos_grafico_categorias(usuario, almacen=almacen)
    cat_labels = datos_cat.get("labels", [])
    cat_data = datos_cat.get("data", [])

    if cat_labels and cat_data:
        # Tabla de datos fuente ubicada en columnas H e I para alimentar el gráfico sin saturar la vista principal
        col_cat_label = 8
        col_cat_val = 9
        fila_datos = r_cat

        c_h1 = ws1.cell(row=fila_datos, column=col_cat_label, value="Categoría")
        c_h1.font = FONT_CABECERA
        c_h1.fill = FILL_CABECERA
        c_h1.border = BORDER_THIN
        c_h1.alignment = ALIGN_CENTER

        c_h2 = ws1.cell(row=fila_datos, column=col_cat_val, value="Productos")
        c_h2.font = FONT_CABECERA
        c_h2.fill = FILL_CABECERA
        c_h2.border = BORDER_THIN
        c_h2.alignment = ALIGN_CENTER

        start_cat_row = fila_datos + 1
        for i, (nombre_cat, cant_prod) in enumerate(zip(cat_labels, cat_data)):
            r_curr = start_cat_row + i
            c_lbl = ws1.cell(row=r_curr, column=col_cat_label, value=nombre_cat)
            c_lbl.font = FONT_DATOS
            c_lbl.fill = FILL_FILA_PAR if (i % 2 == 0) else FILL_FILA_IMPAR
            c_lbl.border = BORDER_THIN
            c_lbl.alignment = ALIGN_LEFT

            c_val = ws1.cell(row=r_curr, column=col_cat_val, value=cant_prod)
            c_val.font = FONT_DATOS_BOLD
            c_val.fill = FILL_FILA_PAR if (i % 2 == 0) else FILL_FILA_IMPAR
            c_val.border = BORDER_THIN
            c_val.alignment = ALIGN_RIGHT
            c_val.number_format = "#,##0"

        end_cat_row = start_cat_row + len(cat_labels) - 1

        # Elemento visual único: Gráfico Donut corporativo centrado en columnas A-F
        chart_cat = DoughnutChart()
        chart_cat.title = "DISTRIBUCIÓN DE ARTÍCULOS POR CATEGORÍA"
        chart_cat.style = 10
        chart_cat.width = 18
        chart_cat.height = 12
        chart_cat.holeSize = 50

        chart_cat.dataLabels = DataLabelList()
        chart_cat.dataLabels.showPercent = True
        chart_cat.dataLabels.showVal = True

        data_ref = Reference(ws1, min_col=col_cat_val, min_row=fila_datos, max_row=end_cat_row)
        labels_ref = Reference(ws1, min_col=col_cat_label, min_row=start_cat_row, max_row=end_cat_row)
        chart_cat.add_data(data_ref, titles_from_data=True)
        chart_cat.set_categories(labels_ref)

        ws1.add_chart(chart_cat, f"A{r_cat}")

    autoajustar_columnas(ws1)

    # ─────────────────────────────────────────────────────────────────────────
    # PESTAÑA 2: EXISTENCIAS Y STOCK DETALLADO
    # ─────────────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="Existencias y Stock")
    crear_banner_hoja(ws2, "AUDACIA — CATÁLOGO DE EXISTENCIAS", f"CORTE: {ahora_str} | ÁMBITO: {nombre_almacen}", 8)

    cabeceras_ex = [
        "Código Interno",
        "Nombre del Producto",
        "Categoría",
        "Almacén",
        "Stock Disponible",
        "Stock en Préstamo",
        "Stock Mínimo",
        "Estado del Stock",
    ]
    r_ex = 4
    aplicar_estilo_fila(ws2, r_ex, cabeceras_ex, es_cabecera=True)
    ws2.freeze_panes = "A5"
    r_ex += 1

    alm_id_filtro = almacen.id if almacen else None
    existencias_qs = obtener_reporte_existencias(usuario, almacen_id=alm_id_filtro)
    start_ex_row = r_ex

    for idx, ex in enumerate(existencias_qs):
        estado_desc = "Óptimo"
        if ex.stock_disponible == 0:
            estado_desc = "AGOTADO / CRÍTICO"
        elif ex.stock_disponible < ex.producto.stock_minimo:
            estado_desc = "BAJO MÍNIMO"

        fila_ex = [
            ex.producto.codigo_interno,
            ex.producto.nombre,
            ex.producto.categoria.nombre if ex.producto.categoria else "Sin categoría",
            ex.almacen.nombre,
            ex.stock_disponible,
            ex.stock_en_prestamo,
            ex.producto.stock_minimo,
            estado_desc,
        ]
        aplicar_estilo_fila(
            ws2,
            r_ex,
            fila_ex,
            es_par=(idx % 2 == 0),
            alineaciones=["left", "left", "left", "left", "right", "right", "right", "center"],
            formatos=["", "", "", "", "#,##0", "#,##0", "#,##0", ""],
        )
        r_ex += 1

    if existencias_qs.exists():
        end_ex_row = r_ex - 1
        totales_ex = [
            "TOTALES",
            "",
            "",
            "",
            f"=SUM(E{start_ex_row}:E{end_ex_row})",
            f"=SUM(F{start_ex_row}:F{end_ex_row})",
            f"=SUM(G{start_ex_row}:G{end_ex_row})",
            "",
        ]
        aplicar_estilo_fila(
            ws2,
            r_ex,
            totales_ex,
            es_total=True,
            alineaciones=["left", "", "", "", "right", "right", "right", ""],
            formatos=["", "", "", "", "#,##0", "#,##0", "#,##0", ""],
        )

    autoajustar_columnas(ws2)

    # ─────────────────────────────────────────────────────────────────────────
    # PESTAÑA 3: ÚLTIMOS MOVIMIENTOS (KARDEX)
    # ─────────────────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="Últimos Movimientos")
    crear_banner_hoja(ws3, "AUDACIA — REGISTRO HISTÓRICO DE MOVIMIENTOS", f"CORTE: {ahora_str} | ÁMBITO: {nombre_almacen}", 9)

    cabeceras_mov = [
        "Folio / ID",
        "Fecha y Hora",
        "Tipo Movimiento",
        "Código Producto",
        "Producto",
        "Almacén Origen",
        "Almacén Destino",
        "Cantidad",
        "Responsable",
    ]
    r_mov = 4
    aplicar_estilo_fila(ws3, r_mov, cabeceras_mov, es_cabecera=True)
    ws3.freeze_panes = "A5"
    r_mov += 1

    movimientos_qs = obtener_reporte_movimientos_periodo(usuario)
    if almacen:
        from django.db.models import Q
        movimientos_qs = movimientos_qs.filter(Q(almacen_origen=almacen) | Q(almacen_destino=almacen))
    movimientos_qs = movimientos_qs[:100]  # Limitar a los 100 más representativos para el dashboard

    start_mov_row = r_mov
    for idx, m in enumerate(movimientos_qs):
        fila_mov = [
            str(m.id)[:8].upper(),
            m.fecha_hora.strftime("%Y-%m-%d %H:%M"),
            m.get_tipo_display(),
            m.producto.codigo_interno,
            m.producto.nombre,
            m.almacen_origen.nombre if m.almacen_origen else "-",
            m.almacen_destino.nombre if m.almacen_destino else "-",
            m.cantidad,
            m.usuario.nombre_completo or m.usuario.username if m.usuario else "Sistema",
        ]
        aplicar_estilo_fila(
            ws3,
            r_mov,
            fila_mov,
            es_par=(idx % 2 == 0),
            alineaciones=["center", "center", "center", "left", "left", "left", "left", "right", "left"],
            formatos=["", "", "", "", "", "", "", "#,##0", ""],
        )
        r_mov += 1

    if movimientos_qs:
        end_mov_row = r_mov - 1
        totales_mov = [
            "TOTAL MOVIMIENTOS",
            "",
            "",
            "",
            "",
            "",
            "",
            f"=SUM(H{start_mov_row}:H{end_mov_row})",
            "",
        ]
        aplicar_estilo_fila(
            ws3,
            r_mov,
            totales_mov,
            es_total=True,
            alineaciones=["left", "", "", "", "", "", "", "right", ""],
            formatos=["", "", "", "", "", "", "", "#,##0", ""],
        )

    autoajustar_columnas(ws3)

    # ─────────────────────────────────────────────────────────────────────────
    # PESTAÑA 4: PRÉSTAMOS PENDIENTES
    # ─────────────────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="Préstamos Pendientes")
    crear_banner_hoja(ws4, "AUDACIA — PRÉSTAMOS ACTIVOS Y EN MORA", f"CORTE: {ahora_str} | ÁMBITO: {nombre_almacen}", 9)

    cabeceras_pres = [
        "Folio",
        "Fecha de Préstamo",
        "Destinatario / Responsable",
        "Código Producto",
        "Producto",
        "Almacén",
        "Original",
        "Pendiente",
        "Estado",
    ]
    r_pres = 4
    aplicar_estilo_fila(ws4, r_pres, cabeceras_pres, es_cabecera=True)
    ws4.freeze_panes = "A5"
    r_pres += 1

    prestamos_qs = obtener_reporte_prestamos_activos(usuario)
    if almacen:
        prestamos_qs = [p for p in prestamos_qs if p.almacen_id == almacen.id]

    start_pres_row = r_pres
    for idx, p in enumerate(prestamos_qs):
        fila_pres = [
            str(p.id)[:8].upper(),
            p.fecha_creacion.strftime("%Y-%m-%d %H:%M"),
            p.destinatario,
            p.producto.codigo_interno,
            p.producto.nombre,
            p.almacen.nombre,
            p.cantidad_original,
            p.cantidad_pendiente_devolucion,
            p.get_estado_display(),
        ]
        aplicar_estilo_fila(
            ws4,
            r_pres,
            fila_pres,
            es_par=(idx % 2 == 0),
            alineaciones=["center", "center", "left", "left", "left", "left", "right", "right", "center"],
            formatos=["", "", "", "", "", "", "#,##0", "#,##0", ""],
        )
        r_pres += 1

    if prestamos_qs:
        end_pres_row = r_pres - 1
        totales_pres = [
            "TOTALES",
            "",
            "",
            "",
            "",
            "",
            f"=SUM(G{start_pres_row}:G{end_pres_row})",
            f"=SUM(H{start_pres_row}:H{end_pres_row})",
            "",
        ]
        aplicar_estilo_fila(
            ws4,
            r_pres,
            totales_pres,
            es_total=True,
            alineaciones=["left", "", "", "", "", "", "right", "right", ""],
            formatos=["", "", "", "", "", "", "#,##0", "#,##0", ""],
        )

    autoajustar_columnas(ws4)

    # Serializar en buffer en memoria
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


# ═════════════════════════════════════════════════════════════════════════════
# EXPORTACIONES DEDICADAS EN EXCEL (EXISTENCIAS, MOVIMIENTOS, PRÉSTAMOS)
# ═════════════════════════════════════════════════════════════════════════════

def exportar_existencias_excel(existencias_qs) -> BytesIO:
    """Exporta el reporte de existencias a un archivo Excel formateado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Existencias"
    ahora_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    crear_banner_hoja(ws, "AUDACIA — CATÁLOGO DE EXISTENCIAS Y STOCK", f"GENERADO: {ahora_str}", 8)
    cabeceras = ["Código Interno", "Producto", "Categoría", "Almacén", "Stock Disp.", "En Préstamo", "Mínimo", "Estado"]
    aplicar_estilo_fila(ws, 4, cabeceras, es_cabecera=True)
    ws.freeze_panes = "A5"

    r = 5
    for idx, ex in enumerate(existencias_qs):
        estado_desc = "Óptimo" if ex.stock_disponible >= ex.producto.stock_minimo else ("Agotado" if ex.stock_disponible == 0 else "Bajo Mínimo")
        fila = [
            ex.producto.codigo_interno,
            ex.producto.nombre,
            ex.producto.categoria.nombre if ex.producto.categoria else "Sin categoría",
            ex.almacen.nombre,
            ex.stock_disponible,
            ex.stock_en_prestamo,
            ex.producto.stock_minimo,
            estado_desc,
        ]
        aplicar_estilo_fila(ws, r, fila, es_par=(idx % 2 == 0), alineaciones=["left", "left", "left", "left", "right", "right", "right", "center"], formatos=["", "", "", "", "#,##0", "#,##0", "#,##0", ""])
        r += 1

    if existencias_qs.exists():
        totales = ["TOTALES", "", "", "", f"=SUM(E5:E{r-1})", f"=SUM(F5:F{r-1})", f"=SUM(G5:G{r-1})", ""]
        aplicar_estilo_fila(ws, r, totales, es_total=True, alineaciones=["left", "", "", "", "right", "right", "right", ""], formatos=["", "", "", "", "#,##0", "#,##0", "#,##0", ""])

    autoajustar_columnas(ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def exportar_movimientos_excel(movimientos_qs) -> BytesIO:
    """Exporta el historial de movimientos a un archivo Excel formateado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kardex de Movimientos"
    ahora_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    crear_banner_hoja(ws, "AUDACIA — REGISTRO HISTÓRICO DE MOVIMIENTOS", f"GENERADO: {ahora_str}", 9)
    cabeceras = ["Folio", "Fecha", "Tipo", "Código", "Producto", "Almacén Origen", "Almacén Destino", "Cantidad", "Responsable"]
    aplicar_estilo_fila(ws, 4, cabeceras, es_cabecera=True)
    ws.freeze_panes = "A5"

    r = 5
    for idx, m in enumerate(movimientos_qs):
        fila = [
            str(m.id)[:8].upper(),
            m.fecha_hora.strftime("%Y-%m-%d %H:%M"),
            m.get_tipo_display(),
            m.producto.codigo_interno,
            m.producto.nombre,
            m.almacen_origen.nombre if m.almacen_origen else "-",
            m.almacen_destino.nombre if m.almacen_destino else "-",
            m.cantidad,
            m.usuario.nombre_completo or m.usuario.username if m.usuario else "Sistema",
        ]
        aplicar_estilo_fila(ws, r, fila, es_par=(idx % 2 == 0), alineaciones=["center", "center", "center", "left", "left", "left", "left", "right", "left"], formatos=["", "", "", "", "", "", "", "#,##0", ""])
        r += 1

    if movimientos_qs.exists():
        totales = ["TOTAL GENERAL", "", "", "", "", "", "", f"=SUM(H5:H{r-1})", ""]
        aplicar_estilo_fila(ws, r, totales, es_total=True, alineaciones=["left", "", "", "", "", "", "", "right", ""], formatos=["", "", "", "", "", "", "", "#,##0", ""])

    autoajustar_columnas(ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def exportar_prestamos_excel(prestamos_qs) -> BytesIO:
    """Exporta el reporte de préstamos a un archivo Excel formateado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Préstamos"
    ahora_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    crear_banner_hoja(ws, "AUDACIA — REPORTE DE PRÉSTAMOS", f"GENERADO: {ahora_str}", 9)
    cabeceras = ["Folio", "Fecha", "Destinatario", "Código", "Producto", "Almacén", "Cant. Original", "Pendiente", "Estado"]
    aplicar_estilo_fila(ws, 4, cabeceras, es_cabecera=True)
    ws.freeze_panes = "A5"

    r = 5
    for idx, p in enumerate(prestamos_qs):
        fila = [
            str(p.id)[:8].upper(),
            p.fecha_creacion.strftime("%Y-%m-%d %H:%M"),
            p.destinatario,
            p.producto.codigo_interno,
            p.producto.nombre,
            p.almacen.nombre,
            p.cantidad_original,
            p.cantidad_pendiente_devolucion,
            p.get_estado_display(),
        ]
        aplicar_estilo_fila(ws, r, fila, es_par=(idx % 2 == 0), alineaciones=["center", "center", "left", "left", "left", "left", "right", "right", "center"], formatos=["", "", "", "", "", "", "#,##0", "#,##0", ""])
        r += 1

    if prestamos_qs:
        totales = ["TOTALES", "", "", "", "", "", f"=SUM(G5:G{r-1})", f"=SUM(H5:H{r-1})", ""]
        aplicar_estilo_fila(ws, r, totales, es_total=True, alineaciones=["left", "", "", "", "", "", "right", "right", ""], formatos=["", "", "", "", "", "", "#,##0", "#,##0", ""])

    autoajustar_columnas(ws)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
